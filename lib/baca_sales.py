# Membaca Excel data sales (sales/daftar-sales.xlsx) lalu mencetak JSON untuk lib/sales.js.
#   python lib/baca_sales.py <file.xlsx>
import json
import sys

from openpyxl import load_workbook

YA = {'ya', 'y', 'yes', 'true', '1', 'aktif'}


def rows(ws):
    """Baris data sebagai dict {judul_kolom_kecil: nilai}, baris kosong dilewati."""
    it = ws.iter_rows(values_only=True)
    head = [str(h or '').strip().lower() for h in next(it)]
    for r in it:
        if not any(v not in (None, '') for v in r):
            continue
        yield {head[i]: r[i] for i in range(min(len(head), len(r)))}


def teks(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def ya(v):
    return teks(v).lower() in YA


def main():
    wb = load_workbook(sys.argv[1], data_only=True)
    out = {'sales': [], 'video': [], 'pengaturan': {}}
    for r in rows(wb['Sales']):
        nama, telp = teks(r.get('nama')), teks(r.get('no. whatsapp'))
        if nama or telp:
            out['sales'].append({'nama': nama, 'telp': telp, 'aktif': ya(r.get('aktif'))})
    for r in rows(wb['Video']):
        folder = teks(r.get('folder'))
        if folder:
            out['video'].append({'kode': teks(r.get('kode')), 'folder': folder, 'aktif': ya(r.get('versi sales')),
                                 '9x16': ya(r.get('9:16')), '16x9': ya(r.get('16:9'))})
    if 'Pengaturan' in wb.sheetnames:
        for r in rows(wb['Pengaturan']):
            k = teks(r.get('kunci'))
            if k:
                out['pengaturan'][k] = teks(r.get('nilai'))
    sys.stdout.buffer.write(json.dumps(out, ensure_ascii=False).encode('utf-8'))


if __name__ == '__main__':
    main()
