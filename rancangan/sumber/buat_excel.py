# Generator Excel rancangan iklan & video Privasimu.
# Pakai: cd motion/rancangan/sumber && python buat_excel.py  → motion/rancangan/Rancangan-Iklan-Video-Privasimu.xlsx
# Isi data ada di data_rancangan.py (video modul, flow unggulan, video pendek, screenshot) dan data_meme.py (48 konsep meme, sound, VO, TTS, kalender).
import datetime as dt
import os
import re
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import data_frasa as DFR
import data_gaya as DG
import data_meme as DM
import data_rancangan as DR

HERE = os.path.dirname(os.path.abspath(__file__))
MOTION = os.path.normpath(os.path.join(HERE, '..', '..'))
OUT = os.path.join(MOTION, 'rancangan', 'Rancangan-Iklan-Video-Privasimu.xlsx')

NAVY, BLUE, GREY = '0B1B4D', '2F6BFF', 'F6F8FC'
thin = Side(style='thin', color='D5DBE7')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HFONT = Font(bold=True, color='FFFFFF', size=11)
HFILL = PatternFill('solid', fgColor=NAVY)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
FILL = {k: PatternFill('solid', fgColor=v) for k, v in dict(hook='FFF3C4', cta='E3F2FF', warn='FDE2E2', ok='E6F6EC', mid='FFF4D6', mod='DCE6FF', zebra='FAFBFD').items()}

# SFX & efek VO yang benar-benar tersedia di pipeline (dibaca dari kode, supaya status di Excel jujur)
audio_js = open(os.path.join(MOTION, 'lib', 'audio.js'), encoding='utf-8').read()
SFX_OK = set(re.findall(r'^\s{2}(\w+)\(bus', audio_js, re.M))
VOFX_OK = os.path.exists(os.path.join(MOTION, 'lib', 'vofx.js'))
DEMO_VO = os.path.exists(os.path.join(MOTION, 'out', '_demo', 'demo-efek-vo.mp3'))
DEMO_SFX = os.path.exists(os.path.join(MOTION, 'out', '_demo', 'demo-sfx-meme.mp3'))

wb = Workbook()
wb.remove(wb.active)


def header(ws, row, cols, height=34):
    for c, t in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=t)
        cell.font = HFONT; cell.fill = HFILL; cell.alignment = CENTER; cell.border = BORDER
    ws.row_dimensions[row].height = height


def widths(ws, ws_widths):
    for c, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def put_row(ws, row, values, center_cols=(), height=None, fill=None, bold_cols=()):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.alignment = CENTER if c in center_cols else WRAP
        cell.border = BORDER
        if fill: cell.fill = fill
        if c in bold_cols: cell.font = Font(bold=True, color=NAVY)
    if height: ws.row_dimensions[row].height = height


def title(ws, text, sub=None, span=8):
    ws['A1'] = text; ws['A1'].font = Font(bold=True, size=16, color=NAVY)
    if sub:
        ws['A2'] = sub; ws['A2'].alignment = WRAP; ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
        ws.row_dimensions[2].height = 46


# ====================================================================== Ringkasan
wr = wb.create_sheet('Ringkasan')
# ====================================================================== Flow Semua Modul (unggulan)
wf = wb.create_sheet('Flow Semua Modul')
F = DR.FLAG
title(wf, 'Flow unggulan: ' + F['judul'], f"Durasi {F['durasi']} · Jenis hook: {F['jenis_hook']} · Hook: {F['hook']}\nCatatan: {F['catatan']}", span=10)
header(wf, 4, ['#', 'Waktu', 'Bab', 'Visual', 'Screenshot (path dari D:\\AI\\privasimu)', 'Crop / sorot / blur', 'Teks layar', 'VO', 'SFX / musik', 'Masuk versi (dtk)'])
for i, (wkt, bab, vis, ss, crop, teks, vo, sfx, versi) in enumerate(F['scenes'], start=1):
    r = 4 + i
    fill = FILL['hook'] if bab == 'Hook' else FILL['cta'] if bab == 'CTA' else (FILL['zebra'] if i % 2 else None)
    put_row(wf, r, [i, wkt, bab, vis, ss, crop, teks, vo, sfx, versi], center_cols=(1, 2, 10), height=62, fill=fill, bold_cols=(3,))
widths(wf, [4, 11, 20, 30, 38, 36, 30, 46, 16, 14])
wf.freeze_panes = 'D5'
wf.auto_filter.ref = f'A4:J{4 + len(F["scenes"])}'

# ====================================================================== Daftar Video Modul + Flow per Modul
wd = wb.create_sheet('Daftar Video Modul')
wp = wb.create_sheet('Flow per Modul')
title(wp, 'Flow per modul: 1 modul = 1 video', 'Setiap blok = satu video. Kolom Screenshot berisi path relatif dari D:\\AI\\privasimu; kolom Crop/sorot/blur berisi koordinat [x, y, lebar, tinggi] pada gambar sumber. "BELUM ADA" = perlu ambil screenshot baru (lihat sheet Screenshot per Modul).', span=10)
header(wp, 4, ['Kode', 'Adegan', 'Waktu', 'Tahap', 'Visual', 'Screenshot (path)', 'Crop / sorot / blur', 'Teks layar', 'VO', 'SFX'])
r = 5
anchor = {}
for m in DR.M:
    anchor[m['kode']] = r
    txt = f"{m['kode']} · {m['modul']} · \"{m['judul']}\" · {m['durasi']} · Hook ({m['jenis_hook']}): {m['hook']}"
    c = wp.cell(row=r, column=1, value=txt); c.font = Font(bold=True, color=NAVY, size=12); c.fill = FILL['mod']; c.alignment = Alignment(wrap_text=True, vertical='center')
    wp.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10); wp.row_dimensions[r].height = 36
    r += 1
    for k, (wkt, tahap, vis, ss, crop, teks, vo, sfx) in enumerate(m['scenes'], start=1):
        fill = FILL['hook'] if tahap in ('Hook', 'Twist') else FILL['cta'] if tahap == 'CTA' else None
        if ss.startswith('BELUM ADA'): fill = FILL['warn']
        put_row(wp, r, [m['kode'], k, wkt, tahap, vis, ss, crop, teks, vo, sfx], center_cols=(1, 2, 3), height=58, fill=fill, bold_cols=(4,))
        r += 1
    r += 1
widths(wp, [7, 7, 12, 11, 30, 40, 40, 30, 46, 16])
wp.freeze_panes = 'E5'

title(wd, 'Daftar video modul (22 video)', 'Ringkasan setiap video modul. Klik kode untuk melompat ke flow-nya. Status screenshot: Siap / Perlu blur / PERLU SCREENSHOT BARU.', span=13)
header(wd, 4, ['Kode', 'Modul', 'Judul', 'Durasi', 'Jenis hook', 'Hook (kalimat pembuka)', 'Target', 'Nada', 'Pesan utama', 'CTA', 'Dasar fakta / klaim', 'Status screenshot', 'Adegan'])
for i, m in enumerate(DR.M, start=5):
    put_row(wd, i, [m['kode'], m['modul'], m['judul'], m['durasi'], m['jenis_hook'], m['hook'], m['target'], m['nada'], m['pesan'], m['cta'], m['fakta'], m['status_ss'], len(m['scenes'])],
            center_cols=(1, 4, 5, 13), height=70, bold_cols=(2,))
    cell = wd.cell(row=i, column=1); cell.hyperlink = f"#'Flow per Modul'!A{anchor[m['kode']]}"; cell.font = Font(bold=True, color=BLUE, underline='single')
    st = wd.cell(row=i, column=12)
    st.fill = FILL['warn'] if 'BARU' in m['status_ss'] else FILL['mid'] if ('blur' in m['status_ss'] or 'crop' in m['status_ss'] or 'perlu' in m['status_ss']) else FILL['ok']
widths(wd, [7, 26, 30, 9, 16, 44, 24, 18, 32, 30, 44, 22, 8])
wd.freeze_panes = 'C5'
wd.auto_filter.ref = f'A4:M{4 + len(DR.M)}'

# ====================================================================== Video Pendek
wv = wb.create_sheet('Video Pendek 5-10-15')
title(wv, 'Video pendek 5 · 10 · 15 detik (meme & serius)', '5 detik = bumper (hook + 1 bukti + logo). 10–15 detik = hook + 1–3 bukti + CTA. Kolom Sumber = video panjang yang bisa dipotong (N = video jadi, M = video modul, A = konsep meme).', span=13)
header(wv, 4, ['Kode', 'Durasi', 'Gaya', 'Jenis hook', 'Judul', 'Hook (0–2 dtk)', 'Isi per detik', 'VO', 'Screenshot (path)', 'SFX', 'CTA', 'Platform', 'Sumber'])
for i, row in enumerate(DR.S, start=5):
    put_row(wv, i, list(row), center_cols=(1, 2, 3, 4), height=56, fill=FILL['zebra'] if i % 2 else None, bold_cols=(5,))
    if row[8].startswith('BELUM') or 'PERLU' in row[8]: wv.cell(row=i, column=9).fill = FILL['warn']
widths(wv, [6, 8, 9, 16, 22, 30, 40, 40, 38, 16, 18, 18, 12])
wv.freeze_panes = 'F5'
wv.auto_filter.ref = f'A4:M{4 + len(DR.S)}'

# ====================================================================== Rancangan Ads Meme (48)
wa = wb.create_sheet('Rancangan Ads Meme')
HEAD = ['ID', 'Judul konsep', 'Format / tren meme', 'Kategori', 'Tujuan utama', 'Target persona', 'Platform & durasi', 'Hook (0–3 detik)', 'Jenis hook',
        'Alur (per detik)', 'Audio meme & SFX', 'Efek VO', 'Modul / layanan', 'CTA', 'Dasar fakta / klaim',
        'Relatable (1–10)', 'Awareness (1–10)', 'Convert (1–10)', 'Share (1–10)', 'SKOR TOTAL', 'Indeks viral', 'Peringkat',
        'Effort produksi', 'Risiko & catatan', 'Status', 'Waktu tayang terbaik', 'kunci']
header(wa, 1, HEAD)
n = len(DM.A)
for i, r0 in enumerate(DM.A, start=2):
    (cid, judul, fmt, kat, tuj, per, plat, hook, alur, aud, vfx, mod, cta, fakta, R, AW, CV, SH, eff, ris, st, wkt) = r0
    row = [cid, judul, fmt, kat, tuj, per, plat, hook, DR.HOOK_A.get(cid, ''), alur, aud, vfx, mod, cta, fakta, R, AW, CV, SH,
           f'=ROUND(Q{i}*W_AW+R{i}*W_CV+S{i}*W_SH,1)', f'=ROUND((P{i}+S{i})/2,1)', f'=RANK(T{i},$T$2:$T${n + 1})', eff, ris, st, wkt, f'=T{i}+ROW()/100000']
    put_row(wa, i, row, center_cols=(1, 9, 16, 17, 18, 19, 20, 21, 22), height=120, fill=FILL['zebra'] if i % 2 == 0 else None, bold_cols=(2,))
    wa.cell(row=i, column=20).font = Font(bold=True, size=12, color=NAVY)
widths(wa, [6, 30, 30, 16, 18, 22, 22, 38, 16, 60, 36, 18, 24, 26, 36, 10, 10, 10, 10, 11, 10, 9, 16, 34, 22, 22, 6])
wa.column_dimensions['AA'].hidden = True
wa.freeze_panes = 'C2'
wa.auto_filter.ref = f'A1:Z{n + 1}'
for col in 'PQRSTU':
    wa.conditional_formatting.add(f'{col}2:{col}{n + 1}', ColorScaleRule(start_type='num', start_value=4, start_color='F8696B', mid_type='num', mid_value=7, mid_color='FFEB84', end_type='num', end_value=10, end_color='63BE7B'))
for c in (16, 17, 18, 19):
    dvs = DataValidation(type='whole', operator='between', formula1='1', formula2='10'); dvs.error = 'Skor 1–10'
    wa.add_data_validation(dvs); dvs.add(f'{get_column_letter(c)}2:{get_column_letter(c)}{n + 1}')
dvh = DataValidation(type='list', formula1='"' + ','.join(DR.HOOK_JENIS) + '"', allow_blank=True)
wa.add_data_validation(dvh); dvh.add(f'I2:I{n + 1}')

# ====================================================================== Screenshot per Modul
wsn = wb.create_sheet('Screenshot per Modul')
title(wsn, 'Inventaris screenshot per modul', 'Semua path relatif dari D:\\AI\\privasimu. Screenshot = tampilan frontend LAMA (boleh dipakai). "Kurasi" = motion/assets/app (sudah crop & blur, siap pakai; dibuat oleh motion/lib/app-shots.js). '
      '"Data demo" = dataroom/03-tenant (lebar 1264): ' + DR.B_DR + '. "UI terbaru" = frontend/tmp/audit4 (lebar 1440, org uji & banyak yang kosong): ' + DR.B_A4 + '.', span=6)
header(wsn, 4, ['Modul', 'Path', 'Sumber', 'Isi', 'Status', 'Catatan crop / blur'])
for i, row in enumerate(DR.INV, start=5):
    put_row(wsn, i, list(row), center_cols=(3, 5), height=38, bold_cols=(1,))
    stt = row[4]
    wsn.cell(row=i, column=5).fill = FILL['ok'] if stt.startswith('Siap') else FILL['warn'] if stt.upper().startswith(('JANGAN', 'KOSONG', 'TIDAK')) else FILL['mid']
last = 4 + len(DR.INV) + 2
wsn.cell(row=last, column=1, value='Screenshot yang perlu diambil baru (data dummy, tanpa nama/merek nyata)').font = Font(bold=True, size=12, color=NAVY)
header(wsn, last + 1, ['Modul', 'Halaman / rute', 'Yang perlu tampil', 'Dipakai di', '', ''])
for k, row in enumerate(DR.BARU, start=last + 2):
    put_row(wsn, k, list(row) + ['', ''], height=40, bold_cols=(1,), fill=FILL['warn'])
widths(wsn, [22, 62, 12, 50, 16, 56])
wsn.freeze_panes = 'B5'
wsn.auto_filter.ref = f'A4:F{4 + len(DR.INV)}'

# ====================================================================== Bank Sound Meme
wsb = wb.create_sheet('Bank Sound Meme')
title(wsb, 'Bank sound meme: hasil cek myinstants.com/en/search/?name=meme (36 tombol, halaman 1)',
      'PENTING: Ketentuan Penggunaan myinstants hanya memberi izin akses untuk penggunaan PRIBADI & NONKOMERSIAL, dan suaranya unggahan pengguna tanpa lisensi dari pemilik asli '
      '(mis. "Faaaah by Taileons"). Jadi file aslinya TIDAK dipakai untuk iklan. Pakai kolom "Versi aman": synth orisinal di motion/lib/audio.js atau rekaman ulang tim di motion/sfx-kustom/.', span=10)
wsb['A2'].fill = FILL['warn']
header(wsb, 4, ['No', 'Nama di myinstants', 'Tautan', 'Jenis', 'Momen di alur iklan', 'Risiko hak / merek', 'Versi aman', "Nama di scenes.js ('sfx')", 'Status versi aman', 'Dipakai di konsep'])
for i, (nm, href, jenis, momen, risk, aman, kode, pakai) in enumerate(DM.SOUND, start=1):
    rr = 4 + i
    names = re.findall(r"'(\w+)'", kode)
    if kode.startswith('sfx-kustom') or 'sfx-kustom' in kode or 'Rekam' in aman:
        status = 'Perlu rekaman tim → motion/sfx-kustom/'
    elif names:
        miss = [x for x in names if x not in SFX_OK]
        status = 'Siap di pipeline' if not miss else 'Perlu dibuat (synth): ' + ', '.join(miss)
    else:
        status = 'Hindari' if 'Hindari' in aman or 'Tidak relevan' in aman else '—'
    put_row(wsb, rr, [i, nm, 'https://www.myinstants.com' + href, jenis, momen, risk, aman, kode, status, pakai], center_cols=(1,), height=46)
    wsb.cell(row=rr, column=3).hyperlink = 'https://www.myinstants.com' + href
    wsb.cell(row=rr, column=3).font = Font(color=BLUE, underline='single')
    wsb.cell(row=rr, column=6).fill = FILL['warn'] if risk.startswith(('Tinggi', 'Brand')) else FILL['mid'] if risk.startswith('Sedang') else FILL['ok']
    wsb.cell(row=rr, column=9).fill = FILL['ok'] if status.startswith('Siap') else FILL['mid'] if status.startswith(('Perlu', 'Rekam')) else PatternFill()
widths(wsb, [5, 28, 34, 26, 30, 30, 40, 26, 30, 30])
wsb.freeze_panes = 'C5'
tips_row = 4 + len(DM.SOUND) + 2
wsb.cell(row=tips_row, column=1, value='Cara rekam suara meme versi tim').font = Font(bold=True, size=12, color=NAVY)
tips = ['Rekam di ruangan senyap, jarak mulut ± 20 cm, 3 take tiap suara (lebay, sedang, singkat). Simpan .wav/.m4a ke motion/sfx-kustom/ (mis. faah.wav).',
        "Panggil di scenes.js: ['w:kata', 'file:faah', 0.8]. Suara otomatis masuk mix & diseimbangkan dengan VO.",
        'Jangan meniru logat/etnis atau persona orang tertentu; cukup ekspresinya (kaget, frustrasi, panik).',
        'Alternatif tanpa rekaman: ElevenLabs v3 dengan tag emosi, mis. "[shouts] FAAAH!" (paket berbayar = lisensi komersial).']
for k, t in enumerate(tips, start=tips_row + 1):
    c = wsb.cell(row=k, column=1, value=f'{k - tips_row}. {t}'); wsb.merge_cells(start_row=k, start_column=1, end_row=k, end_column=10); c.alignment = WRAP; wsb.row_dimensions[k].height = 30
if DEMO_SFX:
    c = wsb.cell(row=tips_row + 6, column=1, value='Dengar semua SFX synth: motion/out/_demo/demo-sfx-meme.mp3'); c.font = Font(bold=True, color=BLUE)

# ====================================================================== Efek VO
wfx = wb.create_sheet('Efek VO')
title(wfx, 'Efek VO (voice effect)',
      ('Tersedia di pipeline: tulis voFx per scene di scenes.js, mis. { id: \'s5\', voFx: \'berat\', ... }. Berlaku untuk TTS maupun VN tim. ' if VOFX_OK else 'Resep efek siap; integrasi ke pipeline (voFx) menyusul. ')
      + ('Dengar semua efek: motion/out/_demo/demo-efek-vo.mp3.' if DEMO_VO else ''), span=6)
header(wfx, 4, ['Preset (voFx)', 'Kesan', 'Kapan dipakai', 'Contoh konsep', 'Resep teknis (ffmpeg)', 'Aman untuk sinkron kata?'])
for i, row in enumerate(DM.FX, start=5):
    put_row(wfx, i, list(row), height=34, bold_cols=(1,))
widths(wfx, [14, 34, 40, 18, 40, 26])

# ====================================================================== Opsi TTS
wt = wb.create_sheet('Opsi TTS')
title(wt, 'Opsi suara (TTS) yang lebih "Indonesia banget"',
      'Cara tercepat tanpa integrasi: buat audio di situs penyedia (mis. ElevenLabs), unduh per bagian, beri nama <KODE>_S1.mp3 dst., taruh di motion/vn/, lalu jalankan node build.js <folder>. '
      'Pipeline menyelaraskan otomatis seperti VN tim. Integrasi API penuh (otomatis & sinkron per kata) bisa ditambahkan setelah ada API key.', span=9)
header(wt, 4, ['Layanan', 'Logat Indonesia', 'Ekspresi / emosi', 'Timing kata (sinkron animasi)', 'Lisensi komersial', 'Kisaran biaya', 'Yang dibutuhkan', 'Catatan / rekomendasi', 'Sumber'])
for i, row in enumerate(DM.TTS, start=5):
    put_row(wt, i, list(row), height=64, bold_cols=(1,))
    if row[8]:
        wt.cell(row=i, column=9).hyperlink = row[8]; wt.cell(row=i, column=9).font = Font(color=BLUE, underline='single')
    if row[0].startswith('ElevenLabs'):
        for c in range(1, 10): wt.cell(row=i, column=c).fill = FILL['ok']
widths(wt, [26, 30, 30, 30, 20, 30, 28, 40, 30])

# ====================================================================== Kalender
wk = wb.create_sheet('Kalender Tayang')
title(wk, 'Kalender tayang: Oktober 2026 → PP 33 berlaku (16 Januari 2027)', 'Video modul (M) + unggulan dijadwalkan berdampingan dengan konten meme (A) dan video pendek (P).', span=7)
VID = ['Unggulan 3 menit (YouTube/LinkedIn) + versi 60 dtk (Reels) · M02', 'M04 · M07', 'M08 · M11', 'M05 · M13', 'M14 · M15', 'M16 · M17', 'M12 · M19', 'M09 · M10',
       'M01 · M03 (3 Des: repost M10)', 'M20 · M21', 'M18 · M22', 'M06 · unggulan versi 30 dtk', 'Rekap: unggulan 60 dtk', 'M02 · M20', 'Unggulan (tayang ulang) · M11', 'Evaluasi & tayang ulang terbaik']
d0 = dt.date(2026, 10, 5)
while d0.weekday() != 0: d0 += dt.timedelta(days=1)
header(wk, 4, ['Minggu', 'Tanggal', 'Fase', 'Momentum', 'Video modul / unggulan', 'Konten organik (meme & pendek)', 'Konten berbayar (konversi)'])
for i, (fase, mom, org, paid) in enumerate(DM.PLAN):
    a = d0 + dt.timedelta(days=7 * i); b = a + dt.timedelta(days=6)
    put_row(wk, 5 + i, [f'M{i + 1}', f'{a.day} {a.strftime("%b")} – {b.day} {b.strftime("%b %Y")}', fase, mom, VID[i] if i < len(VID) else '', org, paid], center_cols=(1,), height=34)
widths(wk, [8, 22, 22, 30, 40, 30, 28])

# ====================================================================== Metodologi Skor
wm = wb.create_sheet('Metodologi Skor')
wm['A1'] = 'Metodologi skor & jenis hook'; wm['A1'].font = Font(bold=True, size=16, color=NAVY)
wm['A2'] = 'Skor adalah ESTIMASI EDITORIAL (penilaian berdasar tren, relevansi, dan kekuatan pesan), bukan data historis. Validasi dengan uji tayang kecil (lihat KPI).'
wm['A2'].alignment = WRAP; wm.merge_cells('A2:F2'); wm.row_dimensions[2].height = 34
wm['A4'] = 'Bobot skor total (ubah angka di kolom C; skor total & peringkat di sheet "Rancangan Ads Meme" ikut berubah)'; wm['A4'].font = Font(bold=True)
for k, (lab, val, nm) in enumerate([('Awareness', 0.35, 'W_AW'), ('Convert', 0.35, 'W_CV'), ('Share', 0.30, 'W_SH')], start=5):
    wm.cell(row=k, column=2, value=lab).font = Font(bold=True)
    c = wm.cell(row=k, column=3, value=val); c.number_format = '0%'; c.fill = PatternFill('solid', fgColor='FFF7CC'); c.border = BORDER
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'Metodologi Skor'!$C${k}")
wm['B8'] = 'Total bobot'; wm['C8'] = '=SUM(C5:C7)'; wm['C8'].number_format = '0%'; wm['B8'].font = Font(bold=True)
crit = [('Relatable', 'Seberapa cepat target merasa "ini gue banget" dalam 3 detik pertama.'),
        ('Awareness', 'Potensi jangkauan: hook kuat, format dikenal luas & masih naik, relevan bagi audiens luas.'),
        ('Convert', 'Kejelasan masalah → modul/layanan, CTA konkret (Pre Check, demo, konsultasi), audiens = pengambil keputusan.'),
        ('Share', 'Dorongan "tag temanmu": humor identitas profesi, mudah ditiru, aman merek.'),
        ('Indeks viral', 'Rata-rata Relatable dan Share.'), ('Skala', '1–3 lemah · 4–6 cukup · 7–8 kuat · 9–10 sangat kuat.')]
wm['A10'] = 'Kriteria skor'; wm['A10'].font = Font(bold=True)
for k, (a, b) in enumerate(crit, start=11):
    wm.cell(row=k, column=2, value=a).font = Font(bold=True)
    c = wm.cell(row=k, column=3, value=b); c.alignment = WRAP; wm.merge_cells(start_row=k, start_column=3, end_row=k, end_column=6); wm.row_dimensions[k].height = 30
HK = [('Reverse psychology', 'Melarang/menantang penonton supaya penasaran.', '"Jangan tonton video ini… kalau perusahaanmu tidak menyimpan data pelanggan." (Unggulan) · "Jangan cek skor kepatuhanmu…" (M02)'),
      ('Relate', 'Situasi yang langsung dikenali target, bikin "ini gue banget".', '"Siapa yang punya RoPA_final_revisi3_FIX.xlsx?" (M04) · "Email paling ditakuti CS" (M07)'),
      ('Anomali', 'Fakta/visual janggal yang menghentikan scroll.', '"Pengguna paling setia aplikasimu mungkin masih kelas 5 SD." (M09) · "12 anak usaha, 12 definisi patuh" (M21)'),
      ('Logika dipatahkan', 'Membantah anggapan umum dengan kalimat pendek.', '"Laporan terbaik bukan yang paling tebal." (M01) · "Kontrak diteken ≠ data aman." (M17)')]
wm['A18'] = 'Jenis hook (wajib ada di setiap video)'; wm['A18'].font = Font(bold=True)
header(wm, 19, ['', 'Jenis', 'Cara kerja', 'Contoh', '', ''])
for k, (a, b, c3) in enumerate(HK, start=20):
    put_row(wm, k, ['', a, b, c3, '', ''], height=40, bold_cols=(2,))
    wm.merge_cells(start_row=k, start_column=4, end_row=k, end_column=6)
wm['A25'] = 'KPI untuk memvalidasi skor (uji 2–3 varian/minggu, dorong yang terbaik dengan anggaran kecil)'; wm['A25'].font = Font(bold=True)
kpi = [('Awareness', '3-second view rate, rata-rata durasi tonton, completion rate, jangkauan unik'), ('Share', 'Share & save per 1.000 view, komentar yang men-tag orang lain'),
       ('Convert', 'CTR ke privasimu.com, Start Pre Check / permintaan demo / konsultasi, biaya per lead'),
       ('Keputusan', 'Awareness & share tertinggi → organik & boost jangkauan; convert tertinggi → iklan berbayar ke pengambil keputusan (LinkedIn/Meta)')]
for k, (a, b) in enumerate(kpi, start=26):
    wm.cell(row=k, column=2, value=a).font = Font(bold=True)
    c = wm.cell(row=k, column=3, value=b); c.alignment = WRAP; wm.merge_cells(start_row=k, start_column=3, end_row=k, end_column=6); wm.row_dimensions[k].height = 30
widths(wm, [4, 20, 34, 40, 20, 20])

# ====================================================================== Katalog Gaya (7 keluarga)
wg = wb.create_sheet('Katalog Gaya')
GH = ['Kode', 'Keluarga', 'Gaya', 'Tampilannya', 'Status', 'Dipakai di', 'Upaya (1–5)', 'Wow (1–10)', 'Cocok merek (1–10)', 'SKOR PRIORITAS', 'Peringkat di keluarga',
      'Gelombang produksi', 'Video usulan', 'Jenis hook', 'Hook (0–3 detik)', 'Alur singkat', 'Meme %', 'Edukasi %', 'Komposisi', 'Modul / fitur',
      'Aset layar (motion/assets/app)', 'Durasi', 'Sumber di plan', 'Teknik (cara buat di mesin kita)', 'Catatan / risiko', 'siap', 'kunci skor', 'kunci peringkat']
header(wg, 1, GH, height=40)
ng = len(DG.G)
GL = ng + 1  # baris terakhir data
for i, g in enumerate(DG.G, start=2):
    meme = None if g['meme'] is None else g['meme'] / 100
    row = [g['kode'], g['keluarga'], g['gaya'], g['tampilan'], g['status'], g['dipakai'], g['upaya'], g['wow'], g['brand'],
           f'=IF(E{i}="Tidak cocok","",ROUND(H{i}*W_GAYA_WOW+I{i}*W_GAYA_MEREK+(11-2*G{i})*W_GAYA_MUDAH,1))',
           f'=IF(Z{i}=1,COUNTIFS($B$2:$B${GL},B{i},$AA$2:$AA${GL},">"&AA{i})+1,"—")',
           g['gelombang'], g['judul'], g['jenis_hook'], g['hook'], g['alur'], meme,
           f'=IF(Q{i}="","",1-Q{i})', f'=IF(Q{i}="","—",TEXT(1-Q{i},"0%")&" edukasi · "&TEXT(Q{i},"0%")&" meme")',
           g['modul'], g['aset'], g['durasi'], g['sumber'], g['teknik'], g['catatan'],
           f'=IF(OR(E{i}="Bisa",E{i}="Berat"),1,0)', f'=IF(Z{i}=1,J{i}-ROW()/100000,"")', f'=B{i}&"|"&K{i}']
    fill = FILL['hook'] if g['gelombang'] == 0 else (FILL['zebra'] if i % 2 == 0 else None)
    put_row(wg, i, row, center_cols=(1, 5, 6, 7, 8, 9, 10, 11, 12, 14, 17, 18, 22), height=78, fill=fill, bold_cols=(3, 13))
    wg.cell(row=i, column=10).font = Font(bold=True, size=12, color=NAVY)
    for c in (17, 18): wg.cell(row=i, column=c).number_format = '0%'
    wg.cell(row=i, column=17).fill = PatternFill('solid', fgColor='FFF7CC')
    st = wg.cell(row=i, column=5)
    st.fill = FILL['ok'] if g['status'] == 'Sudah' else FILL['mid'] if g['status'] == 'Berat' else FILL['warn'] if g['status'] == 'Tidak cocok' else FILL['cta']
widths(wg, [8, 18, 30, 44, 12, 12, 9, 9, 10, 11, 11, 11, 32, 18, 50, 60, 9, 10, 26, 30, 34, 10, 26, 50, 40, 6, 6, 6])
for col in ('Z', 'AA', 'AB'):
    wg.column_dimensions[col].hidden = True
wg.freeze_panes = 'D2'
wg.auto_filter.ref = f'A1:Y{GL}'
for col, lo, mid, hi in (('H', 5, 7, 10), ('I', 5, 7, 10), ('J', 5, 7, 9)):
    wg.conditional_formatting.add(f'{col}2:{col}{GL}', ColorScaleRule(start_type='num', start_value=lo, start_color='F8696B', mid_type='num', mid_value=mid, mid_color='FFEB84', end_type='num', end_value=hi, end_color='63BE7B'))
wg.conditional_formatting.add(f'G2:G{GL}', ColorScaleRule(start_type='num', start_value=1, start_color='63BE7B', mid_type='num', mid_value=3, mid_color='FFEB84', end_type='num', end_value=5, end_color='F8696B'))
for c, a, b in ((7, 1, 5), (8, 1, 10), (9, 1, 10)):
    dv = DataValidation(type='whole', operator='between', formula1=str(a), formula2=str(b)); dv.error = f'Isi {a}–{b}'
    wg.add_data_validation(dv); dv.add(f'{get_column_letter(c)}2:{get_column_letter(c)}{GL}')
dvm = DataValidation(type='decimal', operator='between', formula1='0', formula2='1'); dvm.error = 'Isi 0%–100%'
wg.add_data_validation(dvm); dvm.add(f'Q2:Q{GL}')
dvs = DataValidation(type='list', formula1='"' + ','.join(DG.STATUS) + '"', allow_blank=False)
wg.add_data_validation(dvs); dvs.add(f'E2:E{GL}')

# ====================================================================== Ringkasan Gaya
wq = wb.create_sheet('Ringkasan Gaya')
title(wq, 'Peta gaya motion graphic: 7 keluarga',
      f'{ng} gaya, masing-masing dengan usulan video (hook, alur, modul, komposisi edukasi vs meme). Urutan produksi: one-line art dulu (permintaan), lalu SEMUA tipografi, '
      'lalu SEMUA grafis & bentuk, baru keluarga lain. Daftar lengkap di sheet "Katalog Gaya" (saring kolom Keluarga, urutkan kolom Peringkat).', span=10)
header(wq, 4, ['Kode', 'Keluarga', 'Ciri', 'Gelombang', 'Jumlah gaya', 'Sudah dipakai', 'Bisa', 'Berat', 'Tidak cocok', 'Rata-rata wow'])
KG = "'Katalog Gaya'"
for k, (kode, nama, gel, ciri) in enumerate(DG.KELUARGA, start=5):
    put_row(wq, k, [kode, nama, ciri, gel, f'=COUNTIF({KG}!$B$2:$B${GL},B{k})'] +
            [f'=COUNTIFS({KG}!$B$2:$B${GL},$B{k},{KG}!$E$2:$E${GL},"{s}")' for s in DG.STATUS] +
            [f'=ROUND(AVERAGEIF({KG}!$B$2:$B${GL},B{k},{KG}!$H$2:$H${GL}),1)'], center_cols=(1, 4, 5, 6, 7, 8, 9, 10), height=30, bold_cols=(2,))
rt = 5 + len(DG.KELUARGA)
put_row(wq, rt, ['', 'TOTAL', '', ''] + [f'=SUM({get_column_letter(c)}5:{get_column_letter(c)}{rt - 1})' for c in range(5, 10)] +
        [f'=ROUND(AVERAGE({KG}!$H$2:$H${GL}),1)'], center_cols=(5, 6, 7, 8, 9, 10), height=24, bold_cols=(2, 5, 6, 7, 8, 9, 10), fill=FILL['mod'])

rw = rt + 2
wq.cell(row=rw, column=1, value='Bobot skor prioritas (ubah angka di kolom C; skor & peringkat di "Katalog Gaya" ikut berubah)').font = Font(bold=True, size=12, color=NAVY)
for k, (lab, val, nm, ket) in enumerate([('Wow (daya tarik visual)', 0.45, 'W_GAYA_WOW', 'Seberapa kuat gaya ini menghentikan scroll.'),
                                         ('Cocok merek', 0.35, 'W_GAYA_MEREK', 'Seberapa pas untuk merek B2B yang serius.'),
                                         ('Mudah dibuat', 0.20, 'W_GAYA_MUDAH', 'Dari kolom Upaya: upaya 1 = nilai 9, upaya 5 = nilai 1 (rumus 11 − 2 × upaya).')], start=rw + 1):
    wq.cell(row=k, column=2, value=lab).font = Font(bold=True)
    c = wq.cell(row=k, column=3, value=val); c.number_format = '0%'; c.fill = PatternFill('solid', fgColor='FFF7CC'); c.border = BORDER; c.alignment = Alignment(horizontal='left')
    d = wq.cell(row=k, column=4, value=ket); d.alignment = WRAP; wq.merge_cells(start_row=k, start_column=4, end_row=k, end_column=10)
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'Ringkasan Gaya'!$C${k}")
wq.cell(row=rw + 4, column=2, value='Total bobot').font = Font(bold=True)
c = wq.cell(row=rw + 4, column=3, value=f'=SUM(C{rw + 1}:C{rw + 3})'); c.number_format = '0%'; c.alignment = Alignment(horizontal='left')

rl = rw + 6
wq.cell(row=rl, column=1, value='Cara membaca').font = Font(bold=True, size=12, color=NAVY)
LEG = [('Status', 'Sudah = sudah dipakai di video yang ada · Bisa = bisa dibuat dengan mesin sekarang · Berat = bisa, tapi butuh waktu/aset lebih · Tidak cocok = butuh footage asli, 3D, ilustrator, atau izin.'),
       ('Upaya', '1 = beberapa jam kerja · 3 = sekelas video T01–T03 · 5 = paling berat (banyak objek atau 3D semu).'),
       ('Komposisi', 'Porsi durasi untuk hiburan/format meme vs penjelasan & bukti produk. Ubah kolom "Meme %"; kolom Edukasi & Komposisi mengikuti. 0% = murni edukasi, 100% = full meme.'),
       ('Gelombang', '0 = dikerjakan pertama (one-line art) · 1 = tipografi · 2 = grafis & bentuk · 3–7 = keluarga lain. Di dalam satu gelombang, ikuti kolom Peringkat.'),
       ('Skor', 'Skor wow, cocok merek, dan upaya adalah ESTIMASI EDITORIAL, bukan data. Validasi dengan uji tayang.'),
       ('Aturan isi', 'Fitur hanya dari fakta_produk.json/screenshot asli · angka di hook yang bukan fakta = ilustrasi (diberi label) · "pihak ketiga", bukan "vendor" · tanpa wajah/suara/nama tokoh nyata dan tanpa merek pihak lain.')]
for k, (a, b) in enumerate(LEG, start=rl + 1):
    wq.cell(row=k, column=2, value=a).font = Font(bold=True)
    c = wq.cell(row=k, column=3, value=b); c.alignment = WRAP; wq.merge_cells(start_row=k, start_column=3, end_row=k, end_column=10); wq.row_dimensions[k].height = 34

rp = rl + len(LEG) + 2
wq.cell(row=rp, column=1, value='Urutan produksi: 10 teratas tiap gelombang prioritas (mengikuti skor; berubah otomatis bila bobot/skor diubah)').font = Font(bold=True, size=12, color=NAVY)
header(wq, rp + 1, ['#', 'Tipografi: gaya', 'Video usulan', 'Komposisi', 'Skor', '#', 'Grafis & bentuk: gaya', 'Video usulan', 'Komposisi', 'Skor'])
ambil = lambda kol, nama, n: f'=IFERROR(INDEX({KG}!${kol}$2:${kol}${GL},MATCH("{nama}|"&{n},{KG}!$AB$2:$AB${GL},0)),"")'
for n in range(1, 11):
    k = rp + 1 + n
    put_row(wq, k, [n, ambil('C', 'Tipografi', n), ambil('M', 'Tipografi', n), ambil('S', 'Tipografi', n), ambil('J', 'Tipografi', n),
                    n, ambil('C', 'Grafis & bentuk', n), ambil('M', 'Grafis & bentuk', n), ambil('S', 'Grafis & bentuk', n), ambil('J', 'Grafis & bentuk', n)],
            center_cols=(1, 5, 6, 10), height=30, bold_cols=(2, 7))
satu = next(g for g in DG.G if g['gelombang'] == 0)
c = wq.cell(row=rp + 13, column=1, value=f"Dikerjakan pertama: {satu['kode']} · {satu['gaya']} · video “{satu['judul']}” ({satu['durasi']})"
            + (' — sudah dirakit, tinggal render.' if satu['status'] == 'Sudah' else '.'))
c.font = Font(bold=True, color=NAVY); c.fill = FILL['hook']; wq.merge_cells(start_row=rp + 13, start_column=1, end_row=rp + 13, end_column=10)
widths(wq, [7, 30, 34, 26, 12, 13, 30, 34, 26, 13])
wq.freeze_panes = 'A5'

# ====================================================================== Hook Frasa Viral
wh = wb.create_sheet('Hook Frasa Viral')
title(wh, 'Hook frasa viral (termasuk frasa politik): pemancing perhatian, versi aman',
      'Frasa dipakai sebagai PEMANCING di 0–3 detik pertama, lalu video kembali ke pesan produk. Yang dipakai hanya frasanya (teks + suara tim); '
      'wajah, suara asli, nama, klip video, dan gambar meme tokoh TIDAK dipakai. Konteks tiap frasa cepat berubah: cek ulang sebelum tayang.', span=9)
wh['A2'].fill = FILL['warn']
wh.cell(row=4, column=1, value='Aturan pakai').font = Font(bold=True, size=12, color=NAVY)
for k, t in enumerate(DFR.ATURAN, start=5):
    c = wh.cell(row=k, column=1, value=f'{k - 4}. {t}'); wh.merge_cells(start_row=k, start_column=1, end_row=k, end_column=9); c.alignment = WRAP; wh.row_dimensions[k].height = 30
hf = 5 + len(DFR.ATURAN) + 1
header(wh, hf, ['No', 'Frasa', 'Jenis', 'Konteks singkat', 'Contoh hook Privasimu', 'Modul', 'Cara pakai aman', 'Risiko', 'Catatan'])
for i, row in enumerate(DFR.FRASA, start=1):
    rr = hf + i
    put_row(wh, rr, [i] + list(row), center_cols=(1, 3, 8), height=62, bold_cols=(2,))
    wh.cell(row=rr, column=8).fill = FILL['warn'] if row[6] == 'Tinggi' else FILL['mid'] if row[6] == 'Sedang' else FILL['ok']
widths(wh, [5, 30, 16, 40, 56, 22, 30, 10, 40])
wh.freeze_panes = f'C{hf + 1}'
wh.auto_filter.ref = f'A{hf}:I{hf + len(DFR.FRASA)}'

# ====================================================================== isi Ringkasan
wr['A1'] = 'Rancangan Iklan & Video Privasimu: Q4 2026 → PP 33 berlaku (16 Jan 2027)'; wr['A1'].font = Font(bold=True, size=18, color=NAVY)
wr['A2'] = (f'Isi: 1 flow unggulan "semua modul" (± 3 menit, dengan versi 60 & 30 dtk) · {len(DR.M)} video modul · {len(DR.S)} video pendek 5/10/15 dtk · {len(DM.A)} konsep iklan meme · '
            f'{ng} gaya motion graphic (7 keluarga) · {len(DFR.FRASA)} frasa hook. '
            'Setiap video punya hook (reverse psychology / relate / anomali / logika dipatahkan) dan lokasi screenshot per adegan. Dibuat 29 Sep 2026, diperbarui 30 Sep 2026. '
            'Yang sudah jadi video: N01–N11, M01, M02, A22, T01–T03, GB05 (lihat kolom Status di Katalog Gaya).')
wr['A2'].alignment = WRAP; wr.merge_cells('A2:F2'); wr.row_dimensions[2].height = 46
header(wr, 4, ['Sheet', 'Isi', '', '', '', ''])
TOC = [('Ringkasan Gaya', f'Peta {ng} gaya motion graphic dalam 7 keluarga: jumlah per status, bobot skor, cara membaca, dan 10 teratas tipografi & grafis-bentuk.'),
       ('Katalog Gaya', 'Daftar lengkap gaya: tampilan, status, skor, video usulan (hook, alur, komposisi edukasi vs meme), modul, aset layar, teknik, catatan.'),
       ('Hook Frasa Viral', f'{len(DFR.FRASA)} frasa pemancing (politik, warganet, kantor, suara meme) + aturan pakai aman dan tingkat risikonya.'),
       ('Flow Semua Modul', 'Flow unggulan "Perjalanan Satu Data": semua modul dalam satu video (± 3 menit), per adegan + screenshot. Kolom terakhir menandai adegan untuk versi 60 & 30 dtk.'),
       ('Daftar Video Modul', f'Ringkasan {len(DR.M)} video modul: judul, hook & jenisnya, target, pesan, CTA, dasar klaim, status screenshot.'),
       ('Flow per Modul', 'Flow adegan per video modul: waktu, visual, path screenshot, crop/sorot/blur, teks layar, VO, SFX.'),
       ('Video Pendek 5-10-15', f'{len(DR.S)} video pendek (bumper 5 dtk, 10 dtk, 15 dtk), meme & serius.'),
       ('Rancangan Ads Meme', f'{len(DM.A)} konsep iklan meme dengan skor Awareness · Convert · Share (bobot bisa diubah).'),
       ('Screenshot per Modul', 'Inventaris screenshot: path, isi, status, catatan blur; daftar screenshot yang perlu diambil baru.'),
       ('Bank Sound Meme', 'Hasil cek myinstants + versi aman tiap suara (synth / rekaman tim).'),
       ('Efek VO', 'Preset efek suara VO (berat, tupai, telepon, toa, robot, gema, bisik, dll.).'),
       ('Opsi TTS', 'ElevenLabs, Prosa.ai, Google Chirp 3 HD, Azure: perbandingan & cara pakai.'),
       ('Kalender Tayang', 'Jadwal mingguan Okt 2026 – Jan 2027.'),
       ('Metodologi Skor', 'Bobot skor, kriteria, 4 jenis hook, KPI validasi.')]
for k, (sh, isi) in enumerate(TOC, start=5):
    c = wr.cell(row=k, column=1, value=sh); c.hyperlink = f"#'{sh}'!A1"; c.font = Font(bold=True, color=BLUE, underline='single'); c.border = BORDER
    d = wr.cell(row=k, column=2, value=isi); d.alignment = WRAP; d.border = BORDER; wr.merge_cells(start_row=k, start_column=2, end_row=k, end_column=6); wr.row_dimensions[k].height = 30
r = 5 + len(TOC) + 1
wr.cell(row=r, column=1, value='Urutan produksi yang disarankan').font = Font(bold=True, size=12, color=NAVY)
ORD = ['1. Flow unggulan "Perjalanan Satu Data" (3 menit) + potongan 60 & 30 dtk: fondasi kampanye.',
       '2. Video modul 10 produk inti di landing page: M02 GAP, M04 RoPA, M05 DPIA, M07 DSR, M08 Consent & Cookie, M13 Pihak Ketiga, M14 Transfer Lintas Negara, M11 Insiden, M15 Data Discovery, M19 Priva AI.',
       '3. Video pendek dari potongan N11 & modul (P01–P32) untuk ritme harian.',
       '4. Modul lain (M01, M03, M06, M09, M10, M12, M16–M18, M20–M22), dengan mengambil screenshot baru dulu bila statusnya "PERLU SCREENSHOT BARU".']
for k, t in enumerate(ORD, start=r + 1):
    c = wr.cell(row=k, column=1, value=t); wr.merge_cells(start_row=k, start_column=1, end_row=k, end_column=6); c.alignment = WRAP; wr.row_dimensions[k].height = 30
r = r + len(ORD) + 2
wr.cell(row=r, column=1, value='Catatan penting').font = Font(bold=True, size=12, color=NAVY)
NOTES = ['Screenshot yang dirujuk adalah tampilan frontend LAMA (sesuai arahan: boleh dipakai). Path relatif dari D:\\AI\\privasimu.',
         'JANGAN pakai: dataroom/02-superadmin/holding-dashboard/* (nama perusahaan nyata), dsr/01-overview (merek bank nyata), tabel cookie (nama layanan & domain nyata), data-discovery/03-system-detail (form kredensial), bagian "DPO / Team" di ropa/03 (email).',
         'myinstants.com: hanya untuk pemakaian pribadi & nonkomersial → file aslinya tidak dipakai. "FAAAH" dkk. direkam ulang tim atau dibuat dengan ElevenLabs (paket berbayar).',
         'Suara TTS sekarang (edge-tts) hanya placeholder. Untuk logat Indonesia natural: ElevenLabs, Prosa.ai, atau Google Chirp 3 HD (sheet Opsi TTS).',
         'Klaim produk mengacu fakta_produk.json, katalog modul, dan tampilan aplikasi (kolom "Dasar fakta"). Tokoh (mis. Rina) & angka lucu adalah ilustrasi. Promo "Start Pre Check (gratis)" & "konsultasi gratis": cek masih berlaku.']
for k, t in enumerate(NOTES, start=r + 1):
    c = wr.cell(row=k, column=1, value='• ' + t); wr.merge_cells(start_row=k, start_column=1, end_row=k, end_column=6); c.alignment = WRAP; wr.row_dimensions[k].height = 40
r = r + len(NOTES) + 2
wr.cell(row=r, column=1, value='Top 10 konsep meme (skor total, bobot awal 35/35/30)').font = Font(bold=True, size=12, color=NAVY)
header(wr, r + 1, ['#', 'ID', 'Judul', 'Jenis hook', 'Skor', 'Status'])
tot = lambda x: round(x[15] * .35 + x[16] * .35 + x[17] * .30, 1)
for i, x in enumerate(sorted(DM.A, key=lambda x: (-tot(x), x[0]))[:10], start=1):
    put_row(wr, r + 1 + i, [i, x[0], x[1], DR.HOOK_A.get(x[0], ''), tot(x), x[20]], center_cols=(1, 2, 5), height=22)
widths(wr, [24, 8, 46, 18, 9, 30])

# sheet peta gaya ditaruh tepat setelah Ringkasan
DEPAN = ['Ringkasan', 'Ringkasan Gaya', 'Katalog Gaya', 'Hook Frasa Viral']
wb._sheets = [wb[n] for n in DEPAN] + [ws for ws in wb._sheets if ws.title not in DEPAN]

wb.calculation.fullCalcOnLoad = True
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('OK', OUT)
print(f'gaya {ng} · frasa {len(DFR.FRASA)}')
print(f'video modul {len(DR.M)} ({sum(len(m["scenes"]) for m in DR.M)} adegan) · unggulan {len(DR.FLAG["scenes"])} adegan · pendek {len(DR.S)} · meme {len(DM.A)} · inventaris {len(DR.INV)}')
print('SFX siap:', len(SFX_OK), '· voFx:', VOFX_OK, '· demo VO:', DEMO_VO, '· demo SFX:', DEMO_SFX)
