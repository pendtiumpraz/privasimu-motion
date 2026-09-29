# Membuat Excel data sales untuk versi video per sales: sales/daftar-sales.xlsx
#   python lib/buat_template_sales.py          (tidak menimpa file yang sudah ada)
#   python lib/buat_template_sales.py --timpa  (buat ulang dari nol)
# File berisi nama & nomor HP (data pribadi): TIDAK ikut repo git (lihat .gitignore), jangan dibagikan sembarangan.
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'sales', 'daftar-sales.xlsx')

NAVY, BIRU, KUNING = '0B1B4D', '2F6BFF', 'FFF6C8'
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
ISI_FILL = PatternFill('solid', fgColor=KUNING)

# (kode, folder, judul, versi sales, 9:16, 16:9, catatan)
VIDEO = [
    ('N01', 'ad2-pp33', 'Hitung mundur PP 33/2026', 'Ya', 'Ya', 'Tidak', ''),
    ('N02', 'ad1-awareness', 'Awareness UU PDP', 'Ya', 'Ya', 'Tidak', ''),
    ('N03', 'n03-insiden', 'Insiden 3x24 jam', 'Ya', 'Ya', 'Tidak', ''),
    ('N04', 'n04-ropa', 'RoPA (parodi iklan obat)', 'Ya', 'Ya', 'Tidak', ''),
    ('N05', 'n05-dsr', 'Hak subjek data (DSR)', 'Ya', 'Ya', 'Tidak', ''),
    ('N06', 'n06-anak', 'Data anak & inklusif', 'Ya', 'Ya', 'Tidak', 'kontak dua baris: email tetap, nomor diganti'),
    ('N07', 'n07-konsultan', 'Konsultan PDP', 'Ya', 'Ya', 'Tidak', ''),
    ('N08', 'n08-siap-pp33', 'Program siap PP 33', 'Ya', 'Ya', 'Tidak', ''),
    ('N09', 'n09-ppdp', 'PPDP baru (arcade)', 'Ya', 'Ya', 'Tidak', ''),
    ('N10', 'n10-holding', 'Holding & anak usaha', 'Ya', 'Ya', 'Tidak', ''),
    ('N11', 'n11-meme-dpo', 'Meme DPO', 'Ya', 'Ya', 'Tidak', ''),
    ('M01', 'm01-dasbor', 'Dasbor Kepatuhan', 'Ya', 'Ya', 'Tidak', ''),
    ('M02', 'm02-gap', 'GAP Assessment', 'Ya', 'Ya', 'Tidak', ''),
    ('A22', 'a22-rekap-dpo', 'Rekap Tahunan DPO', 'Tidak', 'Ya', 'Tidak', 'CTA-nya tanpa nomor kontak (konten untuk dibagikan)'),
    ('T01', 't01-stomp', 'Stomp: Satu Platform', 'Ya', 'Ya', 'Tidak', ''),
    ('T02', 't02-tipografi', 'Full Typography: Jangan Tonton', 'Ya', 'Ya', 'Tidak', ''),
    ('T03', 't03-stop-motion', 'Stop Motion: Rating Kebiasaan', 'Ya', 'Ya', 'Tidak', ''),
]

PETUNJUK = [
    'VERSI VIDEO PER SALES',
    '',
    'Tiap video diberi nama & nomor WhatsApp sales di bagian penutup (CTA), menggantikan kontak umum',
    '"support@privasimu.com · 0851 8318 2722". Suara dan isi video lain tetap sama.',
    '',
    'CARA PAKAI',
    '1. Sheet "Sales": isi Nama dan No. WhatsApp, lalu ubah kolom Aktif menjadi Ya. Baris Aktif = Tidak dilewati.',
    '   Nomor boleh ditulis 0812-3456-7890, 081234567890, +62 812 3456 7890, atau 6281234567890.',
    '2. Sheet "Video": pilih video yang dibuatkan versi sales (Ya/Tidak) dan formatnya (9:16 untuk WhatsApp/Reels/TikTok,',
    '   16:9 untuk YouTube/LinkedIn/presentasi).',
    '3. Sheet "Pengaturan" (opsional): teks kontak di layar dan pola nama file.',
    '4. Simpan file ini, lalu klik dua kali RENDER-SALES.bat di folder motion.',
    '',
    'HASIL',
    'Di folder video yang sama, mis. out\\t01-stomp\\t01-stomp-9x16-081234567890.mp4 (nomor ada di nama file).',
    'Video dasar harus sudah dirender dulu (RENDER.bat). Versi yang sudah jadi dan tidak berubah dilewati otomatis,',
    'jadi menambah sales baru hanya merender sales baru itu.',
    '',
    'WAKTU',
    'Hanya bagian akhir video yang dirender ulang: kira-kira 1-2 menit per video per format per sales.',
    'Contoh: 10 sales x 16 video x 1 format = 160 versi = kira-kira 3-5 jam. Kurangi video/format bila perlu.',
    '',
    'PRIVASI',
    'File ini berisi data pribadi (nama & nomor HP). File ini tidak ikut repo git; jangan diunggah ke tempat publik.',
    'Pastikan sales setuju nomornya ditampilkan di materi promosi.',
]


def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'


def main():
    if os.path.exists(OUT) and '--timpa' not in sys.argv:
        print(f'Sudah ada: {OUT} (pakai --timpa untuk membuat ulang)')
        return
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = 'Petunjuk'
    for line in PETUNJUK:
        ws.append([line])
    ws.column_dimensions['A'].width = 120
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    for r in (6, 14, 19, 23):
        ws.cell(row=r, column=1).font = Font(bold=True, color=BIRU)

    ws = wb.create_sheet('Sales')
    header(ws, ['No', 'Nama', 'No. WhatsApp', 'Aktif', 'Catatan'], [6, 30, 22, 10, 50])
    ws.append([1, 'Contoh Sales A', '0812-0000-0001', 'Tidak', 'contoh - ganti dengan data asli lalu ubah Aktif = Ya'])
    ws.append([2, 'Contoh Sales B', '+62 812 0000 0002', 'Tidak', 'format +62 juga boleh'])
    for r in range(4, 104):
        ws.cell(row=r, column=1, value=r - 1)
    dv = DataValidation(type='list', formula1='"Ya,Tidak"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('D2:D200')
    for r in range(2, 104):
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).fill = ISI_FILL

    ws = wb.create_sheet('Video')
    header(ws, ['Kode', 'Folder', 'Judul', 'Versi sales', '9:16', '16:9', 'Catatan'], [7, 18, 32, 12, 8, 8, 52])
    for row in VIDEO:
        ws.append(list(row))
    dv2 = DataValidation(type='list', formula1='"Ya,Tidak"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f'D2:F{len(VIDEO) + 1}')
    for r in range(2, len(VIDEO) + 2):
        for c in (4, 5, 6):
            ws.cell(row=r, column=c).fill = ISI_FILL

    ws = wb.create_sheet('Pengaturan')
    header(ws, ['Kunci', 'Nilai', 'Keterangan'], [22, 34, 80])
    ws.append(['teks_penuh', '{nama} · WA {telp}', 'Menggantikan baris "support@privasimu.com · 0851 8318 2722". Kode: {nama} {telp}'])
    ws.append(['teks_nomor', 'WA {telp}', 'Menggantikan nomor umum yang berdiri sendiri (mis. N06)'])
    ws.append(['pola_nama_file', '{video}-{format}-{telp}', 'Tanpa .mp4. Kode: {video} {format} {telp} {nama}'])
    ws.append(['maks_huruf_nama', 22, 'Nama yang lebih panjang dipotong agar muat di layar'])
    for r in range(2, 6):
        ws.cell(row=r, column=2).fill = ISI_FILL

    wb.save(OUT)
    print(f'Dibuat: {OUT}')


if __name__ == '__main__':
    main()
