# Cek rumus workbook rancangan tanpa Excel/LibreOffice: hitung semua rumus dengan pustaka `formulas`, lalu laporkan
# sel yang galat (#NAME?, #VALUE!, #REF!, #DIV/0!, #N/A) dan beberapa uji kewajaran untuk sheet gaya.
#   pip install formulas
#   python cek_rumus.py            (memeriksa ../Rancangan-Iklan-Video-Privasimu.xlsx)
import os
import sys

import formulas
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FILE = os.path.normpath(os.path.join(HERE, '..', 'Rancangan-Iklan-Video-Privasimu.xlsx'))
NAMA = os.path.basename(FILE).upper()


def main():
    wb = load_workbook(FILE)
    rumus = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    rumus[(ws.title.upper(), c.coordinate)] = c.value
    print(f'rumus di workbook: {len(rumus)}')

    xl = formulas.ExcelModel().loads(FILE).finish()
    sol = xl.calculate()
    nilai = {}
    for k, v in sol.items():
        # kunci: '[FILE]SHEET'!A1 atau rentang
        if '!' not in k or ':' in k.split('!')[-1]:
            continue
        sheet = k.split(']')[1].split("'")[0]
        try:
            val = v.value[0, 0]
        except Exception:
            continue
        nilai[(sheet, k.split('!')[-1])] = val

    galat, kosong = [], []
    for key, f in rumus.items():
        if key not in nilai:
            kosong.append(key)
            continue
        v = nilai[key]
        if isinstance(v, str) and v.startswith('#') or type(v).__name__ == 'XlError':
            galat.append((key, f, v))
    print(f'terhitung: {len(rumus) - len(kosong)} · tidak terhitung: {len(kosong)} · galat: {len(galat)}')
    for key, f, v in galat[:30]:
        print('  GALAT', key, v, '<-', f[:110])
    for key in kosong[:10]:
        print('  TIDAK TERHITUNG', key, rumus[key][:110])

    # ---- uji kewajaran sheet gaya
    KG = 'KATALOG GAYA'
    ws = wb['Katalog Gaya']
    n = ws.max_row
    per = {}
    for i in range(2, n + 1):
        kel, st = ws.cell(row=i, column=2).value, ws.cell(row=i, column=5).value
        rk = nilai.get((KG, f'K{i}'))
        skor = nilai.get((KG, f'J{i}'))
        if st in ('Bisa', 'Berat'):
            per.setdefault(kel, []).append((rk, skor, ws.cell(row=i, column=1).value))
    ok = True
    for kel, rows in per.items():
        ranks = sorted(int(r[0]) for r in rows)
        if ranks != list(range(1, len(rows) + 1)):
            ok = False
            print(f'  PERINGKAT TIDAK UNIK di {kel}: {ranks[:12]}…')
        top = sorted(rows, key=lambda r: int(r[0]))[:3]
        print(f'  {kel}: {len(rows)} gaya siap produksi · 3 teratas: ' + ', '.join(f'{k} ({float(s):.1f})' for _, s, k in top))
    RG = 'RINGKASAN GAYA'
    wq = wb['Ringkasan Gaya']
    for i in range(1, wq.max_row + 1):
        if wq.cell(row=i, column=2).value == 'TOTAL':
            tot = nilai.get((RG, f'E{i}'))
            print(f'  total gaya menurut rumus: {tot} (baris data: {n - 1})')
            ok = ok and int(tot) == n - 1
    contoh = [(k, v) for k, v in nilai.items() if k[0] == RG and k[1][0] in 'BG' and isinstance(v, str) and v and not v.startswith('=')]
    print('  contoh isi 10 teratas:', [v for _, v in contoh[:4]])
    print('  contoh komposisi:', nilai.get((KG, 'S2')), '|', nilai.get((KG, 'S3')))
    sys.exit(1 if (galat or kosong or not ok) else 0)


if __name__ == '__main__':
    main()
